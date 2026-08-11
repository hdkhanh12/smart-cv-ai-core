from __future__ import annotations

from pathlib import Path

import pymupdf
from docx import Document

from ai_core.errors import WarningCode
from ai_core.extraction import CompatibilityFullDocumentProvider
from ai_core.parsers import parse_document
from ai_core.pipeline import process_document
from ai_core.validation import validate_input


def _font_path() -> Path:
    candidates = [
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/Library/Fonts/Arial Unicode.ttf"),
    ]
    for path in candidates:
        if path.exists():
            return path
    raise AssertionError("A Unicode TrueType font is required for PDF parser tests.")


def _write_pdf(path: Path, page_blocks: list[list[tuple[float, float, str]]]) -> Path:
    pdf = pymupdf.open()
    font = _font_path()
    for blocks in page_blocks:
        page = pdf.new_page()
        page.insert_font(fontname="cvfont", fontfile=str(font))
        for x, y, text in blocks:
            page.insert_text((x, y), text, fontname="cvfont", fontsize=10)
    pdf.set_metadata({"author": "Synthetic Test", "title": "Parser Fixture"})
    pdf.save(path)
    pdf.close()
    return path


def test_pdf_parser_one_column_vietnamese_and_metadata(tmp_path: Path) -> None:
    path = _write_pdf(
        tmp_path / "vietnamese.pdf",
        [
            [
                (72, 72, "NGUYỄN VĂN AN"),
                (72, 96, "KỸ NĂNG"),
                (72, 120, "Python, FastAPI, xử lý dữ liệu"),
                (72, 144, "KINH NGHIỆM LÀM VIỆC"),
            ]
        ],
    )

    parsed = parse_document(validate_input(path))

    assert parsed.metadata.page_count == 1
    assert parsed.metadata.author == "Synthetic Test"
    assert parsed.metadata.title == "Parser Fixture"
    assert "NGUYỄN VĂN AN" in parsed.normalized_text
    assert "xử lý dữ liệu" in parsed.normalized_text
    assert parsed.pages[0].page_number == 1


def test_pdf_parser_keeps_two_column_content(tmp_path: Path) -> None:
    path = _write_pdf(
        tmp_path / "two-column.pdf",
        [
            [
                (50, 72, "SKILLS"),
                (50, 96, "Python and SQL"),
                (320, 72, "EXPERIENCE"),
                (320, 96, "Data Engineer at Example"),
            ]
        ],
    )
    parsed = parse_document(validate_input(path))
    assert "SKILLS" in parsed.raw_text
    assert "EXPERIENCE" in parsed.raw_text
    assert "Data Engineer at Example" in parsed.normalized_text


def test_pdf_parser_detects_scan_without_text_layer(tmp_path: Path) -> None:
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.draw_rect(pymupdf.Rect(20, 20, 200, 200), fill=(0.2, 0.2, 0.2))
    path = tmp_path / "scan.pdf"
    pdf.save(path)
    pdf.close()

    parsed = parse_document(validate_input(path))
    assert parsed.normalized_text == ""
    assert WarningCode.NO_TEXT_LAYER in {warning.code for warning in parsed.warnings}


def test_pdf_repeated_header_footer_is_reduced(tmp_path: Path) -> None:
    path = _write_pdf(
        tmp_path / "multipage.pdf",
        [
            [(50, 40, "CURRICULUM VITAE"), (50, 100, f"Unique page {index}"), (50, 800, "Footer")]
            for index in range(1, 4)
        ],
    )
    parsed = parse_document(validate_input(path))
    assert "CURRICULUM VITAE" in parsed.raw_text
    assert "CURRICULUM VITAE" not in parsed.normalized_text
    assert "Footer" not in parsed.normalized_text


def test_docx_parser_preserves_paragraph_table_order_and_unicode(tmp_path: Path) -> None:
    docx = Document()
    docx.core_properties.author = "Synthetic Test"
    docx.add_paragraph("NGUYỄN VĂN AN")
    table = docx.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Kỹ năng"
    table.cell(0, 1).text = "Python"
    table.cell(1, 0).text = "Language"
    table.cell(1, 1).text = "English"
    docx.add_paragraph("WORK EXPERIENCE")
    path = tmp_path / "bilingual.docx"
    docx.save(path)

    parsed = parse_document(validate_input(path))

    assert parsed.metadata.author == "Synthetic Test"
    assert parsed.metadata.page_count is None
    assert "Kỹ năng | Python" in parsed.normalized_text
    assert "Language | English" in parsed.normalized_text
    assert parsed.normalized_text.index("NGUYỄN") < parsed.normalized_text.index("Kỹ năng")
    assert parsed.normalized_text.index("Python") < parsed.normalized_text.index("WORK EXPERIENCE")


def test_debug_mode_writes_raw_and_normalized_text_only_when_requested(tmp_path: Path) -> None:
    path = _write_pdf(tmp_path / "debug.pdf", [[(72, 72, "Synthetic CV text for debugging")]])
    debug_dir = tmp_path / "debug"

    result = process_document(
        path,
        debug_dir=debug_dir,
        extraction_provider=CompatibilityFullDocumentProvider(),
    )

    artifact_dir = debug_dir / result.source_id
    assert (artifact_dir / "raw_text.txt").read_text(encoding="utf-8")
    assert (artifact_dir / "normalized_text.txt").read_text(encoding="utf-8")
    assert (artifact_dir / "profile.json").read_text(encoding="utf-8")
